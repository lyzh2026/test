"""合并导出多篇文章为 Word 文档。"""
import io
import os
import re
import zipfile
from datetime import datetime, timezone
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.ai_analysis import AIAnalysis
from app.models.article import Article

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))), "static", "templates", "export_template.docx")


def _inject_update_fields(doc: Document) -> None:
    """让 Word 打开时自动提示更新域（TOC/页码）。"""
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    doc.settings.element.append(update_fields)


def _extract_headers_from_docx(path: str) -> dict[str, bytes]:
    """提取 docx 中所有 header XML 文件。"""
    headers: dict[str, bytes] = {}
    with zipfile.ZipFile(path, 'r') as z:
        for name in z.namelist():
            if name.startswith('word/header') and name.endswith('.xml'):
                headers[name] = z.read(name)
            if name.startswith('word/_rels/header') and name.endswith('.xml.rels'):
                headers[name] = z.read(name)
    return headers


def _replace_headers_in_docx(buf: io.BytesIO, headers: dict[str, bytes]) -> io.BytesIO:
    """将 header XML 回注到 docx（zip 级别替换）。"""
    buf.seek(0)
    out = io.BytesIO()
    with zipfile.ZipFile(buf, 'r') as zin, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in headers:
                zout.writestr(item, headers[item.filename])
            else:
                zout.writestr(item, zin.read(item.filename))
    out.seek(0)
    return out


def _build_docx_context(articles: list[Article], analyses: list[AIAnalysis | None]) -> dict:
    """构造 docxtpl / python-docx 共用的上下文数据。"""
    now = datetime.now(timezone.utc)
    article_items = []
    for article, analysis in zip(articles, analyses):
        meta_parts = []
        if article.publish_date:
            meta_parts.append(f"发布日期：{article.publish_date.isoformat()}")
        if article.source_unit:
            meta_parts.append(f"来源：{article.source_unit}")

        body_paragraphs = []
        if article.raw_content:
            for para_text in re.split(r'\n\n+', article.raw_content):
                stripped = para_text.strip()
                if stripped:
                    body_paragraphs.append(stripped)

        article_items.append({
            'title': article.original_title or '无标题',
            'meta': ' | '.join(meta_parts),
            'summary': analysis.summary if analysis else '',
            'keywords': '、'.join(analysis.keywords) if analysis and analysis.keywords else '',
            'categories': '、'.join(c.get('label', '') for c in (analysis.categories or [])) if analysis else '',
            'body_paragraphs': body_paragraphs,
        })

    return {
        'doc_title': '拾讯文章合集',
        'gen_date': now.strftime('%Y-%m-%d'),
        'article_count': len(articles),
        'articles': article_items,
    }


async def _fetch_articles(article_ids: list[str], session: AsyncSession) -> tuple[list[Article], list[AIAnalysis | None]]:
    """按 ID 顺序批量获取文章及 AI 分析。"""
    articles: list[Article] = []
    analyses: list[AIAnalysis | None] = []
    for aid in article_ids:
        res = await session.execute(
            select(Article, AIAnalysis)
            .join(AIAnalysis, AIAnalysis.article_id == Article.id, isouter=True)
            .where(Article.id == aid)
        )
        row = res.one_or_none()
        if row:
            articles.append(row[0])
            analyses.append(row[1])
    return articles, analyses


def _build_without_template(context: dict) -> io.BytesIO:
    """模式 B：无模板，用 python-docx 从零构建。"""
    doc = Document()

    # 页面边距
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # ── 封面页 ──
    cover_section = doc.sections[0]
    cover_section.different_first_page_header_footer = True

    # 标题
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(context['doc_title'])
    run.font.size = Pt(28)
    run.font.bold = True

    doc.add_paragraph()  # 空行

    # 日期
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run(f"生成日期：{context['gen_date']}")
    run.font.size = Pt(12)

    # 文章数量
    count_para = doc.add_paragraph()
    count_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = count_para.add_run(f"共收录 {context['article_count']} 篇文章")
    run.font.size = Pt(12)

    doc.add_page_break()

    # ── 目录页 ──
    doc.add_heading('目录', level=1)
    toc_para = doc.add_paragraph()
    # TOC 域代码：只列 Heading 1
    run = toc_para.add_run()
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar_begin)

    run2 = toc_para.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.text = 'TOC \\o "1-1" \\h \\z \\u'
    run2._r.append(instrText)

    run3 = toc_para.add_run()
    fldChar_sep = OxmlElement('w:fldChar')
    fldChar_sep.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fldChar_sep)

    run4 = toc_para.add_run('（请按 Ctrl+A 后按 F9 更新目录）')

    run5 = toc_para.add_run()
    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')
    run5._r.append(fldChar_end)

    doc.add_page_break()

    # ── 各文章内容 ──
    for idx, article in enumerate(context['articles']):
        # 标题（Heading 1 → 进入 TOC）
        doc.add_heading(article['title'], level=1)

        # 元信息
        if article['meta']:
            meta_para = doc.add_paragraph()
            meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = meta_para.add_run(article['meta'])
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

        # AI 摘要
        if article['summary']:
            doc.add_heading('AI 摘要', level=2)
            doc.add_paragraph(article['summary'])
            if article['keywords']:
                doc.add_paragraph(f"关键词：{article['keywords']}")

        # 分类
        if article['categories']:
            doc.add_paragraph(f"分类：{article['categories']}")

        # 正文
        if article['body_paragraphs']:
            doc.add_heading('正文', level=2)
            for para in article['body_paragraphs']:
                doc.add_paragraph(para)

        # 分页（最后一篇不需要）
        if idx < len(context['articles']) - 1:
            doc.add_page_break()

    # 页脚页码
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        if footer.paragraphs:
            footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = footer.paragraphs[0].add_run()
            # PAGE 域代码
            fldChar_begin = OxmlElement('w:fldChar')
            fldChar_begin.set(qn('w:fldCharType'), 'begin')
            run._r.append(fldChar_begin)
            run2 = footer.paragraphs[0].add_run()
            instrText = OxmlElement('w:instrText')
            instrText.text = 'PAGE'
            run2._r.append(instrText)
            run3 = footer.paragraphs[0].add_run()
            fldChar_end = OxmlElement('w:fldChar')
            fldChar_end.set(qn('w:fldCharType'), 'end')
            run3._r.append(fldChar_end)

    _inject_update_fields(doc)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _build_with_template(context: dict) -> io.BytesIO | None:
    """模式 A：使用 docxtpl 填充模板。"""
    try:
        from docxtpl import DocxTemplate, RichText
    except ImportError:
        return None

    if not os.path.isfile(TEMPLATE_PATH):
        return None

    # 1. 提取原始 header（保护水印/logo）
    headers = _extract_headers_from_docx(TEMPLATE_PATH)

    # 2. 准备 RichText 内容
    tpl = DocxTemplate(TEMPLATE_PATH)
    articles_rt = []
    for art in context['articles']:
        body_rt = RichText()
        for para in art['body_paragraphs']:
            body_rt.add(para + '\n')
        articles_rt.append({
            'title': art['title'],
            'meta': art['meta'],
            'summary': art['summary'],
            'keywords': art['keywords'],
            'categories': art['categories'],
            'body': body_rt,
        })

    render_ctx = {
        'doc_title': context['doc_title'],
        'gen_date': context['gen_date'],
        'article_count': context['article_count'],
        'articles': articles_rt,
    }

    tpl.render(render_ctx)
    _inject_update_fields(tpl.docx)

    buf = io.BytesIO()
    tpl.save(buf)
    buf.seek(0)

    # 3. 回注 header 保护水印/logo
    if headers:
        buf = _replace_headers_in_docx(buf, headers)

    buf.seek(0)
    return buf


async def export_merged_doc(
    article_ids: list[str],
    use_template: bool,
    session: AsyncSession,
) -> StreamingResponse:
    """入口函数：合并导出多篇文章。"""
    articles, analyses = await _fetch_articles(article_ids, session)
    if not articles:
        from app.utils.response import error
        from fastapi import Request
        return error(2003, "未找到指定的文章", http_status=404, request=Request(scope={"type": "http"}))

    context = _build_docx_context(articles, analyses)

    # 尝试模板模式
    if use_template:
        buf = _build_with_template(context)
        if buf:
            filename = f"拾讯文章合集_{context['gen_date']}.docx"
            encoded = quote(filename, safe="")
            return StreamingResponse(
                buf,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
            )

    # 回退到无模板模式
    buf = _build_without_template(context)
    filename = f"拾讯文章合集_{context['gen_date']}.docx"
    encoded = quote(filename, safe="")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )