"""文章查询 + 统计 + 分类编辑接口（PRD 3.2.3）。"""
import io
import logging
import re
import uuid
from datetime import date, timedelta
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_session
from app.dependencies.auth import current_admin
from app.models.ai_analysis import AIAnalysis
from app.models.article import Article
from app.models.system_config import SystemConfig
from app.utils.response import error, success

logger = logging.getLogger("shixun.articles")

router = APIRouter(prefix="/api/v1", tags=["articles"])


async def _get_category_labels(session: AsyncSession) -> list[str]:
    """从 SystemConfig 读取分类标签，无配置时返回默认值。始终包含"未分类"。"""
    result = await session.execute(
        select(SystemConfig).where(SystemConfig.key == "category_labels")
    )
    cfg = result.scalar_one_or_none()
    if cfg and cfg.value and isinstance(cfg.value, list):
        labels = [str(l).strip() for l in cfg.value if str(l).strip()]
    else:
        labels = [
            "最新政策", "数字经济", "人工智能", "数据要素", "通信",
            "申报", "潜在商机", "具身智能", "车路云协同", "新型工业化", "算力",
        ]
    if "未分类" not in labels:
        labels.append("未分类")
    return labels


def _serialize_article(article: Article, analysis: AIAnalysis | None) -> dict:
    return {
        "id": str(article.id),
        "task_id": str(article.task_id) if article.task_id else None,
        "original_title": article.original_title,
        "source_unit": article.source_unit,
        "original_link": article.original_link,
        "publish_date": article.publish_date.isoformat() if article.publish_date else None,
        "status": article.status,
        "failed_reason": article.failed_reason,
        "bookmarked": article.bookmarked,
        "created_at": article.created_at.isoformat() if article.created_at else None,
        "updated_at": article.updated_at.isoformat() if article.updated_at else None,
        "ai": {
            "categories": (analysis.categories or []) if analysis else [],
            "summary": (analysis.summary or "") if analysis else "",
            "keywords": (analysis.keywords or []) if analysis else [],
            "model_used": analysis.model_used if analysis else None,
            "processing_time_ms": analysis.processing_time_ms if analysis else None,
        },
    }


@router.get("/articles")
async def list_articles(
    request: Request,
    date_from: date | None = Query(default=None),
    date_to: date | None = Query(default=None),
    category: str | None = Query(default=None, max_length=64),
    keyword: str | None = Query(default=None, max_length=128),
    status: str | None = Query(default=None, max_length=32),
    bookmarked: bool | None = Query(default=None),
    task_id: str | None = Query(default=None, max_length=64),
    limit: int = Query(default=20, ge=1, le=100),
    offset: int = Query(default=0, ge=0),
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    stmt = select(Article, AIAnalysis).join(AIAnalysis, AIAnalysis.article_id == Article.id, isouter=True)
    conds = []
    if date_from:
        conds.append(Article.publish_date >= date_from)
    if date_to:
        conds.append(Article.publish_date <= date_to)
    if status:
        conds.append(Article.status == status)
    if bookmarked is not None:
        conds.append(Article.bookmarked == bookmarked)
    if keyword:
        like = f"%{keyword}%"
        conds.append(Article.original_title.ilike(like))
    if task_id:
        try:
            conds.append(Article.task_id == uuid.UUID(task_id))
        except ValueError:
            pass
    if category:
        conds.append(AIAnalysis.categories.op("@>")([{"label": category}]))
    if conds:
        stmt = stmt.where(and_(*conds))
    stmt = stmt.order_by(Article.publish_date.desc(), Article.created_at.desc()).limit(limit).offset(offset)

    res = await session.execute(stmt)
    rows = res.all()
    items = [_serialize_article(a, ai) for (a, ai) in rows]

    count_stmt = select(func.count(Article.id)).join(
        AIAnalysis, AIAnalysis.article_id == Article.id, isouter=True
    )
    if conds:
        count_stmt = count_stmt.where(and_(*conds))
    total = (await session.execute(count_stmt)).scalar_one()

    return success({"items": items, "total": total, "limit": limit, "offset": offset}, request=request)


@router.get("/articles/category-labels")
async def list_category_labels(
    request: Request,
    session: AsyncSession = Depends(get_session),
):
    """获取当前分类标签列表（公开接口，无需登录）。"""
    labels = await _get_category_labels(session)
    return success({"labels": labels}, request=request)


@router.get("/articles/{article_id}")
async def get_article(
    article_id: str,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    res = await session.execute(
        select(Article, AIAnalysis)
        .join(AIAnalysis, AIAnalysis.article_id == Article.id, isouter=True)
        .where(Article.id == article_id)
    )
    row = res.one_or_none()
    if not row:
        return error(2003, "文章不存在", http_status=404, request=request)
    article, analysis = row
    data = _serialize_article(article, analysis)
    data["raw_content"] = article.raw_content
    return success(data, request=request)


@router.get("/articles/{article_id}/export/doc")
async def export_article_doc(
    article_id: str,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """导出单篇文章为 .docx 文件。"""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    res = await session.execute(
        select(Article, AIAnalysis)
        .join(AIAnalysis, AIAnalysis.article_id == Article.id, isouter=True)
        .where(Article.id == article_id)
    )
    row = res.one_or_none()
    if not row:
        return error(2003, "文章不存在", http_status=404, request=Request(scope={"type": "http"}))

    article, analysis = row

    doc = Document()

    # 标题
    title_para = doc.add_heading(article.original_title or "无标题", level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 元信息
    meta_parts = []
    if article.publish_date:
        meta_parts.append(f"发布日期：{article.publish_date.isoformat()}")
    if article.source_unit:
        meta_parts.append(f"来源：{article.source_unit}")
    if article.original_link:
        meta_parts.append(f"原文链接：{article.original_link}")
    if meta_parts:
        meta_para = doc.add_paragraph()
        meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = meta_para.add_run(" | ".join(meta_parts))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    doc.add_paragraph("")  # 空行

    # AI 摘要
    if analysis and analysis.summary:
        doc.add_heading("AI 摘要", level=2)
        doc.add_paragraph(analysis.summary)
        if analysis.keywords:
            doc.add_paragraph("关键词：" + "、".join(analysis.keywords))

    # AI 分类
    if analysis and analysis.categories:
        cats = "、".join(c.get("label", "") for c in analysis.categories)
        doc.add_paragraph(f"分类：{cats}")

    doc.add_paragraph("")  # 空行

    # 正文
    if article.raw_content:
        doc.add_heading("正文", level=2)
        for para_text in article.raw_content.split("\n"):
            stripped = para_text.strip()
            if stripped:
                doc.add_paragraph(stripped)

    # 写入内存
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # 文件名：去掉特殊字符
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', (article.original_title or "article")[:60])
    filename = f"{safe_title}.docx"
    encoded_filename = quote(filename, safe="")

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.patch("/articles/{article_id}/bookmark")
async def toggle_bookmark(
    article_id: str,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    article = await session.get(Article, article_id)
    if not article:
        return error(2003, "文章不存在", http_status=404, request=request)
    article.bookmarked = not article.bookmarked
    await session.commit()
    return success({"bookmarked": article.bookmarked}, request=request)


@router.get("/stats/categories")
async def stats_categories(
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """按 AI 分类标签聚合计数。"""
    analyses = await session.execute(select(AIAnalysis.article_id, AIAnalysis.categories))
    counter: dict[str, int] = {}
    for _, cats in analyses.all():
        for cat in cats:
            label = cat.get("label", "未分类")
            counter[label] = counter.get(label, 0) + 1
    items = [{"label": k, "count": v} for k, v in counter.items()]
    items.sort(key=lambda x: x["count"], reverse=True)
    return success({"items": items}, request=request)


@router.get("/stats/sources")
async def stats_sources(
    request: Request,
    limit: int = Query(default=10, ge=1, le=50),
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """按来源 source_unit 分组计数。"""
    stmt = (
        select(Article.source_unit, func.count(Article.id))
        .where(Article.source_unit.isnot(None))
        .group_by(Article.source_unit)
        .order_by(func.count(Article.id).desc())
        .limit(limit)
    )
    res = await session.execute(stmt)
    items = [{"source": source, "count": count} for source, count in res.all()]
    return success({"items": items}, request=request)


@router.post("/articles/{article_id}/reanalyze")
async def reanalyze_article(
    article_id: str,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """对失败或未处理状态的文章手动触发 AI 重新分析。"""
    article = await session.get(Article, article_id)
    if not article:
        return error(2003, "文章不存在", http_status=404, request=request)
    if article.status not in ("raw", "failed_retryable", "failed_permanent"):
        return error(1001, f"当前状态 ({article.status}) 不允许重新分析", http_status=400, request=request)

    # 改回 raw 状态以便 analyze_article 入口放行
    if article.status != "raw":
        article.status = "raw"
        article.failed_reason = None
        await session.commit()

    from app.modules.ai.service import analyze_article
    try:
        result = await analyze_article(article_id)
        return success({"ok": True, "status": result.get("status")}, request=request)
    except Exception as e:
        logger.error("reanalyze failed: article=%s %r", article_id, e)
        return error(5001, f"重新分析失败: {str(e)}", http_status=500, request=request)


@router.delete("/articles/{article_id}")
async def delete_article(
    article_id: str,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    article = await session.get(Article, article_id)
    if not article:
        return error(2003, "文章不存在", http_status=404, request=request)

    # 先删除关联的 AI 分析记录
    analysis = await session.execute(
        select(AIAnalysis).where(AIAnalysis.article_id == article_id)
    )
    analysis_row = analysis.scalar_one_or_none()
    if analysis_row:
        await session.delete(analysis_row)

    await session.delete(article)
    await session.commit()
    return success({"ok": True}, request=request)


@router.post("/articles/batch-delete")
async def batch_delete_articles(
    payload: dict,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """批量删除文章（幂等，不存在的 ID 自动跳过）。"""
    article_ids = payload.get("article_ids", [])
    if not article_ids or not isinstance(article_ids, list):
        return error(1001, "article_ids 必须为非空数组", http_status=400, request=request)
    if len(article_ids) > 100:
        return error(1001, "单次最多删除 100 篇文章", http_status=400, request=request)

    deleted = 0
    for aid in article_ids:
        article = await session.get(Article, aid)
        if not article:
            continue
        analysis = await session.execute(
            select(AIAnalysis).where(AIAnalysis.article_id == aid)
        )
        analysis_row = analysis.scalar_one_or_none()
        if analysis_row:
            await session.delete(analysis_row)
        await session.delete(article)
        deleted += 1

    await session.commit()
    return success({"deleted": deleted}, request=request)


@router.get("/stats/daily")
async def stats_daily(
    request: Request,
    days: int = Query(default=14, ge=1, le=90),
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    today = date.today()
    start = today - timedelta(days=days - 1)

    counts_stmt = (
        select(Article.publish_date, Article.status, func.count(Article.id))
        .where(Article.publish_date >= start)
        .group_by(Article.publish_date, Article.status)
    )
    res = await session.execute(counts_stmt)
    daily_map: dict[str, dict[str, int]] = {}
    for d, st, c in res.all():
        ds = d.isoformat()
        daily_map.setdefault(ds, {"total": 0})
        daily_map[ds][st] = c
        daily_map[ds]["total"] += c

    daily = []
    for i in range(days):
        d = (start + timedelta(days=i)).isoformat()
        item = daily_map.get(d, {"total": 0})
        item["date"] = d
        daily.append(item)

    return success({"daily": daily, "from": start.isoformat(), "to": today.isoformat()}, request=request)


@router.get("/dates")
async def distinct_dates(
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    res = await session.execute(
        select(Article.publish_date, func.count(Article.id))
        .group_by(Article.publish_date)
        .order_by(Article.publish_date.desc())
        .limit(60)
    )
    items = [{"date": d.isoformat(), "count": c} for d, c in res.all()]
    return success({"items": items}, request=request)


@router.put("/articles/{article_id}/categories")
async def update_article_categories(
    article_id: str,
    payload: dict,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """修改 AI 分类结果（最多 3 个标签）。"""
    categories = payload.get("categories", [])
    if not isinstance(categories, list):
        return error(1001, "categories 必须为数组", http_status=400, request=request)
    if len(categories) > 3:
        return error(1001, "分类标签最多 3 个", http_status=400, request=request)

    valid_labels = set(await _get_category_labels(session) + ["未分类"])
    for c in categories:
        if not isinstance(c, dict) or "label" not in c:
            return error(1001, "每个分类必须包含 label 字段", http_status=400, request=request)
        if c["label"] not in valid_labels:
            return error(1001, f"无效分类标签：{c['label']}", http_status=400, request=request)

    # "未分类"具有排他性：有它就不能有别的标签
    if any(c.get("label") == "未分类" for c in categories):
        categories = [c for c in categories if c.get("label") == "未分类"][:1]

    article = await session.get(Article, article_id)
    if not article:
        return error(2003, "文章不存在", http_status=404, request=request)

    existing = await session.execute(select(AIAnalysis).where(AIAnalysis.article_id == article_id))
    analysis = existing.scalar_one_or_none()
    if analysis:
        analysis.categories = categories
    else:
        analysis = AIAnalysis(article_id=article_id, categories=categories)
        session.add(analysis)

    # 如果文章是失败状态，标记为已处理
    if article.status in ("failed_retryable", "failed_permanent"):
        article.status = "processed"
        article.failed_reason = None

    await session.commit()
    return success({"categories": categories}, request=request)


@router.post("/articles/export/merged-doc")
async def merged_export(
    payload: dict,
    request: Request,
    _: object = Depends(current_admin),
    session: AsyncSession = Depends(get_session),
):
    """合并导出多篇文章为一份 Word 文档。"""
    article_ids = payload.get("article_ids", [])
    use_template = payload.get("use_template", False)

    if not article_ids or not isinstance(article_ids, list):
        return error(1001, "article_ids 必须为非空数组", http_status=400, request=request)
    if len(article_ids) > 100:
        return error(1001, "单次最多导出 100 篇文章", http_status=400, request=request)

    from app.modules.articles.docx_export import export_merged_doc
    return await export_merged_doc(article_ids, use_template, session)

